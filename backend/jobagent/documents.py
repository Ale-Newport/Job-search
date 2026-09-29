"""Evidence-preserving imports and immutable, explicitly approved documents."""

from __future__ import annotations

import difflib
import hashlib
import io
import re
import zipfile
from collections import defaultdict
from functools import lru_cache
from pathlib import Path
from xml.sax.saxutils import escape

from docx import Document as WordDocument
from pypdf import PdfReader
from reportlab.lib import colors
from reportlab.lib.enums import TA_LEFT
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import inch
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer

from .db import decode_row, dumps, now, uid
from .matching import extract_skills

MAX_IMPORT = 20 * 1024 * 1024
SUBMITTED_STATUSES = {
    "APPLYING",
    "APPLIED",
    "CONFIRMED",
    "RECRUITER_SCREEN",
    "ASSESSMENT",
    "TECHNICAL_TEST",
    "INTERVIEW",
    "FINAL_INTERVIEW",
    "OFFER",
    "REJECTED",
    "WITHDRAWN",
    "GHOSTED",
}
MIME_TYPES = {
    ".pdf": "application/pdf",
    ".docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    ".txt": "text/plain",
    ".md": "text/markdown",
    ".csv": "text/csv",
    ".json": "application/json",
}


def extract_text(content: bytes, filename: str) -> tuple[str, str]:
    suffix = Path(filename).suffix.lower()
    if len(content) > MAX_IMPORT:
        raise ValueError("Document exceeds the 20 MB limit")
    if not content:
        raise ValueError("Document is empty")
    try:
        if suffix == ".pdf":
            reader = PdfReader(io.BytesIO(content))
            if reader.is_encrypted:
                raise ValueError("Unlock the PDF before importing it")
            if len(reader.pages) > 100:
                raise ValueError("Document exceeds the 100-page import limit")
            text = "\n".join(page.extract_text() or "" for page in reader.pages)
        elif suffix == ".docx":
            with zipfile.ZipFile(io.BytesIO(content)) as archive:
                if sum(item.file_size for item in archive.infolist()) > 50 * 1024 * 1024:
                    raise ValueError("Expanded DOCX exceeds the safety limit")
            doc = WordDocument(io.BytesIO(content))
            text = "\n".join(
                [p.text for p in doc.paragraphs]
                + [" | ".join(cell.text for cell in row.cells) for table in doc.tables for row in table.rows]
            )
        elif suffix in {".txt", ".md", ".csv", ".json"}:
            text = content.decode("utf-8-sig")
        else:
            raise ValueError("Supported imports: PDF, DOCX, UTF-8 TXT, Markdown, CSV and JSON")
    except ValueError:
        raise
    except Exception as exc:
        raise ValueError("Document could not be read; check its format and encryption") from exc
    text = text.replace("\x00", "").strip()[:500000]
    if not text:
        raise ValueError("No readable text was found. Export a text-based PDF or paste the text as a TXT file")
    return text, MIME_TYPES[suffix]


def propose_facts(text: str, source: str) -> list[dict]:
    """Preserve source lines verbatim; never infer dates, names or qualifications."""
    categories = {
        "education": ("education", "qualifications", "academic", "educación"),
        "experience": ("experience", "employment", "work history", "experiencia"),
        "project": ("projects", "personal projects", "proyectos"),
        "skill": ("skills", "technical skills", "technologies", "habilidades"),
        "publication": ("publications", "research", "publicaciones"),
    }
    proposals = []
    seen = set()

    def add(category, key, value):
        identity = (category, value.strip())
        if value.strip() and identity not in seen:
            seen.add(identity)
            proposals.append(
                {
                    "category": category,
                    "key": key,
                    "value": value.strip(),
                    "source": source,
                    "verification_status": "unverified",
                    "locked": False,
                    "notes": "Extracted from your document. Review context and accuracy before verifying.",
                }
            )

    for email in re.findall(r"[\w.+-]+@[\w.-]+\.[A-Za-z]{2,}", text):
        add("personal", "email", email)
    for url in re.findall(r"https?://[^\s<>]+", text):
        clean = url.rstrip(".,;)")
        key = "github" if "github.com/" in clean else "linkedin" if "linkedin.com/" in clean else "website"
        add("link", key, clean)
    category = "imported"
    for line in text.splitlines():
        line = line.strip()
        if not line:
            continue
        heading = line.lower().strip(":# ")
        heading_category = next((cat for cat, words in categories.items() if heading in words), None)
        if heading_category:
            category = heading_category
            continue
        if len(line) >= 4:
            add(category, f"{category}_{len(proposals) + 1}", line[:10000])
        if len(proposals) >= 300:
            break
    # Explicit skill names found in the document are still proposals, not verified claims.
    for skill in extract_skills(text):
        add("skill", skill.replace(" ", "_"), skill)
    return proposals


def save_version(
    db,
    data_dir: Path,
    document: dict,
    content: bytes,
    filename: str,
    mime_type: str,
    text: str,
    fact_ids=None,
    approved=False,
):
    version_id = uid()
    suffix = Path(filename).suffix.lower()
    target_dir = Path(data_dir) / "documents" / document["id"]
    target_dir.mkdir(parents=True, exist_ok=True)
    path = target_dir / (version_id + suffix)
    # 'xb' prevents overwriting a registered version even under concurrent requests.
    with path.open("xb") as file:
        file.write(content)
    path.chmod(0o600)
    try:
        with db.transaction() as conn:
            row = conn.execute(
                "SELECT COALESCE(MAX(version),0)+1 AS n FROM document_versions WHERE document_id=?", (document["id"],)
            ).fetchone()
            version = row["n"]
            conn.execute(
                "INSERT INTO document_versions VALUES(?,?,?,?,?,?,?,?,?,?,?)",
                (
                    version_id,
                    document["id"],
                    version,
                    str(path.resolve()),
                    Path(filename).name,
                    mime_type,
                    hashlib.sha256(content).hexdigest(),
                    text,
                    dumps(fact_ids or []),
                    int(approved),
                    now(),
                ),
            )
            conn.execute("UPDATE documents SET updated_at=? WHERE id=?", (now(), document["id"]))
    except BaseException:
        path.unlink(missing_ok=True)
        raise
    return decode_row(db.one("SELECT * FROM document_versions WHERE id=?", (version_id,)))


def attach_document(db, conn, application_id, version_id, kind):
    application = conn.execute("SELECT status FROM applications WHERE id=?", (application_id,)).fetchone()
    existing = conn.execute(
        "SELECT document_version_id FROM application_documents WHERE application_id=? AND kind=?",
        (application_id, kind),
    ).fetchall()
    if any(row["document_version_id"] == version_id for row in existing):
        return
    if application["status"] in SUBMITTED_STATUSES:
        raise ValueError("Documents of a submitted application cannot be replaced")
    previous = [row["document_version_id"] for row in existing]
    conn.execute("DELETE FROM application_documents WHERE application_id=? AND kind=?", (application_id, kind))
    conn.execute("INSERT INTO application_documents VALUES(?,?,?)", (application_id, version_id, kind))
    db.event(
        application_id,
        application["status"],
        f"Selected {kind} version {version_id}; previous versions: {', '.join(previous) or 'none'}",
        "manual",
        conn=conn,
    )


def import_document(db, data_dir: Path, content: bytes, filename: str):
    filename = Path(filename).name[:200]
    text, mime = extract_text(content, filename)
    document_id, stamp = uid(), now()
    document = {
        "id": document_id,
        "name": filename,
        "kind": "import",
        "job_id": None,
        "created_at": stamp,
        "updated_at": stamp,
    }
    db.execute("INSERT INTO documents VALUES(?,?,?,?,?,?)", tuple(document.values()))
    try:
        version = save_version(db, data_dir, document, content, filename, mime, text)
    except BaseException:
        db.execute("DELETE FROM documents WHERE id=?", (document_id,))
        raise
    candidate_id = db.one("SELECT id FROM candidates LIMIT 1")["id"]
    proposed = propose_facts(text, f"document:{version['id']}")
    with db.transaction() as conn:
        for fact in proposed:
            fact.update(id=uid(), candidate_id=candidate_id, created_at=stamp, updated_at=stamp)
            conn.execute(
                "INSERT INTO facts(id,candidate_id,category,key,value,source,verification_status,locked,notes,created_at,updated_at) VALUES(?,?,?,?,?,?,?,?,?,?,?)",
                (
                    fact["id"],
                    candidate_id,
                    fact["category"],
                    fact["key"],
                    fact["value"],
                    fact["source"],
                    "unverified",
                    0,
                    fact["notes"],
                    stamp,
                    stamp,
                ),
            )
    return {"document": document, "version": version, "proposed_facts": proposed}


def select_facts(db, fact_ids=None):
    if fact_ids:
        facts = db.query(f"SELECT * FROM facts WHERE id IN ({','.join('?' for _ in fact_ids)})", fact_ids)
        if len(facts) != len(set(fact_ids)):
            raise ValueError("One or more referenced facts do not exist")
        if any(f["verification_status"] != "verified" and not f["locked"] for f in facts):
            raise ValueError("Only verified or locked facts can support a generated document")
    else:
        facts = db.query(
            "SELECT * FROM facts WHERE verification_status='verified' OR locked=1 ORDER BY category,created_at"
        )
    if not facts:
        raise ValueError("Verify candidate facts before generating a document")
    return facts


def compose_document(job: dict, kind: str, facts: list[dict]) -> str:
    relevant = extract_skills(job.get("description", ""))
    facts = sorted(facts, key=lambda f: (-sum(skill in f["value"].lower() for skill in relevant), f["created_at"]))
    name = next((f["value"] for f in facts if f["key"].lower() in {"name", "full_name", "full name"}), "")
    if kind == "cover_letter":
        professional = [
            f for f in facts if f["category"] in {"experience", "project", "education", "skill", "publication"}
        ]
        if not professional:
            raise ValueError("A cover letter needs verified professional facts")
        evidence = "\n\n".join(f["value"] for f in professional[:10])
        return f"Application — {job['title']}\n{job['company']}\n\nDear hiring team,\n\nI am applying for the {job['title']} position at {job['company']}. The following is relevant information from my professional background:\n\n{evidence}\n\nThank you for considering my application. I would welcome the opportunity to discuss the role.\n\n{name}".strip()
    grouped = defaultdict(list)
    for fact in facts:
        if fact["category"] not in {
            "preference",
            "work_authorization",
            "salary",
            "answer",
            "saved_answer",
            "sensitive",
        }:
            grouped[fact["category"]].append(fact["value"])
    sections = [name or "Professional profile", f"Application: {job['title']} · {job['company']}"]
    labels = {
        "personal": "Contact",
        "link": "Links",
        "education": "Education",
        "experience": "Experience",
        "project": "Projects",
        "skill": "Skills",
        "publication": "Publications",
        "imported": "Additional background",
    }
    for category in ("personal", "link", "experience", "education", "project", "skill", "publication", "imported"):
        if grouped.get(category):
            sections.extend(["", labels.get(category, category.title()), *grouped.pop(category)])
    for category, values in grouped.items():
        sections.extend(["", category.replace("_", " ").title(), *values])
    return "\n".join(sections)


@lru_cache(maxsize=1)
def pdf_font():
    # Use system Unicode fonts without redistributing proprietary font files.
    candidates = [
        Path("/System/Library/Fonts/Supplemental/Arial Unicode.ttf"),
        Path("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"),
        Path("/System/Library/Fonts/Supplemental/Arial.ttf"),
    ]
    for path in candidates:
        if path.is_file():
            pdfmetrics.registerFont(TTFont("MeridianUnicode", str(path)))
            return "MeridianUnicode"
    return "Helvetica"


def render_pdf(text: str) -> bytes:
    output = io.BytesIO()
    document = SimpleDocTemplate(
        output,
        pagesize=(595.28, 841.89),
        rightMargin=48,
        leftMargin=48,
        topMargin=46,
        bottomMargin=46,
        title=text.splitlines()[0],
        author="",
    )
    styles = getSampleStyleSheet()
    font = pdf_font()
    styles.add(
        ParagraphStyle(
            name="MeridianBody",
            fontName=font,
            fontSize=10,
            leading=15,
            textColor=colors.HexColor("#21303b"),
            spaceAfter=7,
            alignment=TA_LEFT,
        )
    )
    styles.add(
        ParagraphStyle(
            name="MeridianTitle",
            parent=styles["Heading1"],
            fontName=font,
            fontSize=20,
            leading=25,
            textColor=colors.HexColor("#113f42"),
            spaceAfter=12,
        )
    )
    story = []
    for index, line in enumerate(text.splitlines()):
        if not line.strip():
            story.append(Spacer(1, 0.08 * inch))
        else:
            story.append(Paragraph(escape(line), styles["MeridianTitle"] if index == 0 else styles["MeridianBody"]))
    document.build(story)
    return output.getvalue()


def generate_document(db, data_dir, job_id, kind, fact_ids=None, approved=False):
    if kind not in {"cv", "cover_letter"}:
        raise ValueError("Document kind must be cv or cover_letter")
    job = db.one("SELECT * FROM jobs WHERE id=?", (job_id,))
    if job is None:
        raise ValueError("Job not found")
    facts = select_facts(db, fact_ids)
    text = compose_document(job, kind, facts)
    existing = db.one("SELECT * FROM documents WHERE job_id=? AND kind=? ORDER BY created_at LIMIT 1", (job_id, kind))
    previous = (
        db.one(
            "SELECT text FROM document_versions WHERE document_id=? ORDER BY version DESC LIMIT 1", (existing["id"],)
        )
        if existing
        else None
    )
    if not previous:
        previous = db.one(
            "SELECT text FROM document_versions v JOIN documents d ON d.id=v.document_id WHERE d.kind='import' ORDER BY v.created_at DESC LIMIT 1"
        )
    diff = "\n".join(
        difflib.unified_diff(
            (previous["text"] if previous else "").splitlines(),
            text.splitlines(),
            fromfile="previous",
            tofile="proposed",
            lineterm="",
        )
    )
    response = {
        "text": text,
        "diff": diff,
        "fact_ids": [f["id"] for f in facts],
        "facts_used": [{"id": f["id"], "value": f["value"], "source": f["source"]} for f in facts],
        "approved": bool(approved),
        "kind": kind,
    }
    if approved:
        pdf = render_pdf(text)
        document = existing
        if not document:
            stamp = now()
            document = {
                "id": uid(),
                "name": f"{job['company']} — {job['title']} ({kind})",
                "kind": kind,
                "job_id": job_id,
                "created_at": stamp,
                "updated_at": stamp,
            }
            db.execute("INSERT INTO documents VALUES(?,?,?,?,?,?)", tuple(document.values()))
        response["document"] = document
        response["version"] = save_version(
            db,
            Path(data_dir),
            document,
            pdf,
            f"{kind}-{job_id[:8]}.pdf",
            "application/pdf",
            text,
            response["fact_ids"],
            True,
        )
        application = db.one("SELECT id,status FROM applications WHERE job_id=?", (job_id,))
        if application and application["status"] not in SUBMITTED_STATUSES:
            with db.transaction() as conn:
                attach_document(db, conn, application["id"], response["version"]["id"], kind)
    return response


def document_path(db, data_dir, version_id: str) -> tuple[Path, dict]:
    version = db.one("SELECT * FROM document_versions WHERE id=?", (version_id,))
    if not version:
        raise ValueError("Document version not found")
    path = Path(version["path"]).resolve()
    root = (Path(data_dir) / "documents").resolve()
    if not path.is_relative_to(root) or not path.is_file():
        raise ValueError("Registered document is missing or outside the documents directory")
    if hashlib.sha256(path.read_bytes()).hexdigest() != version["content_hash"]:
        raise ValueError("Document content no longer matches its immutable registered version")
    return path, decode_row(version)
