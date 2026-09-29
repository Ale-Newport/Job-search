"""Anonymized synthetic fixtures only: no private CV contents or identifiers."""

import io
import json

import pytest
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE

from jobagent.cv_parser import extract_docx_text, propose_cv_facts
from jobagent.db import Database
from jobagent.documents import extract_text, import_document, reimport_document


def hyperlink(paragraph, label, target):
    element = OxmlElement("w:hyperlink")
    identifier = paragraph.part.relate_to(target, RELATIONSHIP_TYPE.HYPERLINK, is_external=True)
    element.set(qn("r:id"), identifier)
    run, text = OxmlElement("w:r"), OxmlElement("w:t")
    text.text = label
    run.append(text)
    element.append(run)
    paragraph._p.append(element)


def word_bytes(document):
    stream = io.BytesIO()
    document.save(stream)
    return stream.getvalue()


@pytest.fixture
def cv_bytes():
    doc = Document()
    doc.add_paragraph("MORGAN EXAMPLE")
    contact = doc.add_paragraph("morgan@example.test | +44 (0)7000000000 | ")
    hyperlink(contact, "LinkedIn", "https://www.linkedin.com/in/morgan-example/")
    contact.add_run(" | Bristol, UK")
    doc.add_paragraph("EDUCATION")
    doc.add_paragraph("Example University | MSc Computing | Sep 2098 - Sep 2099")
    doc.add_paragraph("Relevant Modules: Information Retrieval, Systems Verification.", style="List Bullet")
    doc.add_paragraph("")
    doc.add_paragraph("Sample College | BSc Mathematics (First Class) | Sep 2020 - Jun 2023")
    doc.add_paragraph("Relevant Modules: Algebra, Statistics.", style="List Bullet")
    doc.add_paragraph("PROFESSIONAL EXPERIENCE")
    doc.add_paragraph("Data Developer | Example Systems | Jun 2023 - Aug 2024")
    doc.add_paragraph(
        "Processed 73,000+ records, improving review accuracy by 17.5% while reducing latency by 42%.",
        style="List Bullet",
    )
    doc.add_paragraph(
        "Built a retrieval prototype in Python and preserved 99.2% of source identifiers.", style="List Bullet"
    )
    doc.add_paragraph("")
    doc.add_paragraph("Teaching Assistant | Sample College | Sep 2024 - Jun 2025")
    doc.add_paragraph("Supported 64 students in weekly debugging labs.", style="List Bullet")
    doc.add_paragraph("PROJECTS")
    doc.add_paragraph("Grid Navigator")
    doc.add_paragraph("Implemented a simulator handling 21,000 concurrent events.", style="List Bullet")
    doc.add_paragraph("")
    doc.add_paragraph("Open Online Course | Introductory Computing")
    doc.add_paragraph("Completed the introductory course and its exercises.", style="List Bullet")
    doc.add_paragraph("SKILLS")
    doc.add_paragraph("C# | R | Scala | Haskell | Prolog | Python | Fastify | AWS (EC2, S3) | Python | HTML & CSS")
    doc.add_paragraph("LANGUAGES & SOFT SKILLS")
    doc.add_paragraph("English (fluent) | German (Basic).")
    return word_bytes(doc)


def test_full_cv_preserves_sections_metrics_dates_links_and_exact_skills(cv_bytes):
    text, warnings = extract_docx_text(cv_bytes)
    facts, summary = propose_cv_facts(text.strip(), "document:test-version")
    by_category = {}
    for fact in facts:
        by_category.setdefault(fact["category"], []).append(fact)
        assert fact["verification_status"] == "unverified"
        assert fact["locked"] is False
        assert fact["source"] == "document:test-version"
        assert "Extracted line" in fact["notes"]
    assert not warnings
    assert len(by_category["education"]) == 2
    assert len(by_category["experience"]) == 2
    assert len(by_category["project"]) == 2
    assert len(by_category["language"]) == 2
    developer = by_category["experience"][0]["value"]
    assert "Data Developer | Example Systems | Jun 2023 - Aug 2024" in developer
    assert "73,000+" in developer and "17.5%" in developer and "42%" in developer
    assert "99.2% of source identifiers." in developer
    assert developer.count("• ") == 2
    assert "Teaching Assistant" not in developer
    assert "Sep 2098 - Sep 2099" in by_category["education"][0]["value"]
    assert "graduation_year" not in {fact["key"] for fact in facts}
    assert by_category["link"][0]["value"] == "https://www.linkedin.com/in/morgan-example/"
    assert {fact["value"] for fact in by_category["skill"]} == {
        "C#",
        "R",
        "Scala",
        "Haskell",
        "Prolog",
        "Python",
        "Fastify",
        "AWS (EC2, S3)",
        "HTML & CSS",
    }
    assert not any(fact["key"] in {"professional_experience_years", "sponsorship_required"} for fact in facts)
    assert summary["fact_count"] == len(facts)


def test_docx_text_preserves_tables_content_controls_and_body_order():
    doc = Document()
    doc.add_paragraph("EDUCATION")
    table = doc.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Example Academy"
    table.cell(0, 1).text = "2020 - 2024"
    doc.add_paragraph("BSc Computing, Distinction")
    doc.add_paragraph("PROFESSIONAL EXPERIENCE")
    control, content, paragraph, run, text = [
        OxmlElement(tag) for tag in ("w:sdt", "w:sdtContent", "w:p", "w:r", "w:t")
    ]
    text.text = "Engineer | Example Studio | 2024 - 2025"
    run.append(text)
    paragraph.append(run)
    content.append(paragraph)
    control.append(content)
    doc.element.body.insert(-1, control)
    doc.add_paragraph("Implemented 12 deployments.", style="List Bullet")
    text, _ = extract_docx_text(word_bytes(doc))
    assert text.index("EDUCATION") < text.index("Example Academy") < text.index("PROFESSIONAL EXPERIENCE")
    assert text.index("Engineer | Example Studio") < text.index("Implemented 12 deployments.")
    facts, _ = propose_cv_facts(text, "document:tables")
    education = next(fact for fact in facts if fact["category"] == "education")
    assert all(part in education["value"] for part in ("Example Academy", "2020 - 2024", "BSc Computing"))
    experience = next(fact for fact in facts if fact["category"] == "experience")
    assert "Implemented 12 deployments." in experience["value"]


def test_header_contact_hyperlinks_and_field_codes_are_read_without_following():
    doc = Document()
    doc.sections[0].header.paragraphs[0].text = "Morgan Example"
    p = doc.sections[0].header.add_paragraph("Email: morgan@example.test Phone: +1 (555) 010-9876")
    doc.add_paragraph("Portfolio ")
    hyperlink(doc.paragraphs[-1], "Project", "https://example.test/project(v2)")
    p = doc.add_paragraph("Website")
    run = OxmlElement("w:r")
    field = OxmlElement("w:instrText")
    field.text = ' HYPERLINK "https://example.test/hidden-url" '
    run.append(field)
    p._p.append(run)
    doc.add_paragraph("SKILLS")
    doc.add_paragraph("C++")
    text, _ = extract_docx_text(word_bytes(doc))
    assert "https://example.test/project(v2)" in text
    assert "https://example.test/hidden-url" in text
    facts, _ = propose_cv_facts(text, "document:header")
    assert any(fact["key"] == "email" and fact["value"] == "morgan@example.test" for fact in facts)
    assert any(fact["key"] == "phone" and fact["value"] == "+1 (555) 010-9876" for fact in facts)
    assert any(fact["value"] == "https://example.test/project(v2)" for fact in facts)


def test_skill_proposals_do_not_infer_framework_languages_or_short_name_equivalents():
    facts, _ = propose_cv_facts(
        "TECHNICAL SKILLS\nDjango | React | GitHub | C# | R\nLANGUAGES\nSpanish (native)", "document:test"
    )
    assert {fact["value"] for fact in facts if fact["category"] == "skill"} == {"Django", "React", "GitHub", "C#", "R"}
    assert not any(fact["value"] in {"Python", "JavaScript", "Git"} for fact in facts)
    assert next(fact for fact in facts if fact["category"] == "language")["value"] == "Spanish (native)"


def test_project_contact_does_not_become_candidate_identity():
    facts, _ = propose_cv_facts(
        "Morgan Example\nPROJECTS\nDirectory tool\n• Stored support@example.test contact data.", "document:test"
    )
    assert not any(fact["key"] == "email" for fact in facts)
    assert "support@example.test" in next(fact for fact in facts if fact["category"] == "project")["value"]


def test_tracked_deletions_are_omitted_with_warning():
    doc = Document()
    doc.add_paragraph("PROFESSIONAL EXPERIENCE")
    p = doc.add_paragraph("Engineer | Example | 2020 - 2022")
    deletion, run, text = [OxmlElement(tag) for tag in ("w:del", "w:r", "w:delText")]
    text.text = "Invented obsolete claim"
    run.append(text)
    deletion.append(run)
    p._p.append(deletion)
    text, warnings = extract_docx_text(word_bytes(doc))
    assert "Invented obsolete claim" not in text
    assert any("Tracked changes" in warning for warning in warnings)


def test_import_persists_source_linked_unverified_facts_and_immutable_content(tmp_path, cv_bytes):
    db = Database(tmp_path / "parser.sqlite")
    imported = import_document(db, tmp_path, cv_bytes, "anonymous-cv.docx")
    version = imported["version"]
    assert version["approved"] is False or version["approved"] == 0
    assert set(version["fact_ids"]) == {fact["id"] for fact in imported["proposed_facts"]}
    persisted = db.query("SELECT * FROM facts")
    assert all(fact["source"] == f"document:{version['id']}" for fact in persisted)
    assert all(fact["verification_status"] == "unverified" and not fact["locked"] for fact in persisted)
    assert json.loads(db.one("SELECT fact_ids FROM document_versions")["fact_ids"]) == version["fact_ids"]
    from pathlib import Path

    assert Path(version["path"]).read_bytes() == cv_bytes
    assert "PROFESSIONAL EXPERIENCE" in version["text"]


def test_oversized_evidence_fails_without_silent_truncation_or_partial_import(tmp_path):
    db = Database(tmp_path / "parser.sqlite")
    with pytest.raises(ValueError, match="20,000"):
        import_document(db, tmp_path, ("PROJECTS\nLong project\n" + "a" * 20001).encode(), "oversized.txt")
    assert db.one("SELECT COUNT(*) AS n FROM documents")["n"] == 0
    assert db.one("SELECT COUNT(*) AS n FROM facts")["n"] == 0
    with pytest.raises(ValueError, match="500,000"):
        extract_text(b"x" * 500001, "large.txt")


def test_reimport_preserves_prior_version_and_protected_facts_then_is_idempotent(tmp_path, cv_bytes):
    from pathlib import Path

    db = Database(tmp_path / "repair.sqlite")
    old = import_document(db, tmp_path, cv_bytes, "anonymous-cv.docx")
    old_version = db.one("SELECT * FROM document_versions WHERE id=?", (old["version"]["id"],))
    old_file = Path(old_version["path"]).read_bytes()
    db.execute("DELETE FROM settings WHERE key=?", (f"document_parser:{old_version['id']}",))
    verified, locked, manual = [item["id"] for item in old["proposed_facts"][:3]]
    db.execute("UPDATE facts SET verification_status='verified' WHERE id=?", (verified,))
    db.execute("UPDATE facts SET locked=1 WHERE id=?", (locked,))
    db.execute("UPDATE facts SET source='manual' WHERE id=?", (manual,))
    protected = {
        identifier: db.one("SELECT * FROM facts WHERE id=?", (identifier,)) for identifier in (verified, locked, manual)
    }
    other = import_document(db, tmp_path, b"Other Example\nSKILLS\nZig", "other-cv.txt")
    other_ids = {item["id"] for item in other["proposed_facts"]}
    result = reimport_document(db, tmp_path, old["document"]["id"])
    assert result["document"]["id"] == old["document"]["id"]
    assert result["version"]["id"] != old_version["id"]
    assert result["version"]["version"] == 2
    assert result["version"]["path"] != old_version["path"]
    assert Path(result["version"]["path"]).read_bytes() == old_file
    assert db.one("SELECT * FROM document_versions WHERE id=?", (old_version["id"],)) == old_version
    assert Path(old_version["path"]).read_bytes() == old_file
    assert db.one("SELECT COUNT(*) AS n FROM documents")["n"] == 2
    assert len(result["superseded_fact_ids"]) == len(old["proposed_facts"]) - 3
    for identifier in result["superseded_fact_ids"]:
        changed = db.one("SELECT * FROM facts WHERE id=?", (identifier,))
        assert changed["verification_status"] == "rejected"
        assert result["version"]["id"] in changed["notes"]
        history = db.one("SELECT payload FROM fact_history WHERE fact_id=?", (identifier,))
        assert json.loads(history["payload"])["verification_status"] == "unverified"
    for identifier, before in protected.items():
        assert db.one("SELECT * FROM facts WHERE id=?", (identifier,)) == before
    assert all(
        db.one("SELECT verification_status FROM facts WHERE id=?", (identifier,))["verification_status"] == "unverified"
        for identifier in other_ids
    )
    assert all(item["verification_status"] == "unverified" and not item["locked"] for item in result["proposed_facts"])
    retry = reimport_document(db, tmp_path, old["document"]["id"])
    assert retry["already_current"] is True
    assert retry["version"]["id"] == result["version"]["id"]
    assert db.one("SELECT COUNT(*) AS n FROM document_versions WHERE document_id=?", (old["document"]["id"],))["n"] == 2


def test_reimport_rolls_back_new_version_and_rejections_on_failure(tmp_path, cv_bytes):
    from pathlib import Path
    import sqlite3

    db = Database(tmp_path / "failure.sqlite")
    old = import_document(db, tmp_path, cv_bytes, "anonymous-cv.docx")
    db.execute("DELETE FROM settings WHERE key=?", (f"document_parser:{old['version']['id']}",))
    old_facts = db.query("SELECT * FROM facts ORDER BY id")
    db.execute("CREATE TRIGGER fail_new_fact BEFORE INSERT ON facts BEGIN SELECT RAISE(ABORT,'test failure'); END")
    with pytest.raises(sqlite3.IntegrityError, match="test failure"):
        reimport_document(db, tmp_path, old["document"]["id"])
    assert db.one("SELECT COUNT(*) AS n FROM document_versions")["n"] == 1
    assert db.query("SELECT * FROM facts ORDER BY id") == old_facts
    assert len(list(Path(old["version"]["path"]).parent.iterdir())) == 1


def test_reimport_api_uses_existing_document_and_rejects_changed_source(tmp_path, cv_bytes):
    from pathlib import Path
    from fastapi import FastAPI
    from fastapi.testclient import TestClient
    from jobagent.core import router

    db = Database(tmp_path / "api.sqlite")
    old = import_document(db, tmp_path, cv_bytes, "anonymous-cv.docx")
    db.execute("DELETE FROM settings WHERE key=?", (f"document_parser:{old['version']['id']}",))
    app = FastAPI()
    app.state.db, app.state.data_dir = db, tmp_path
    app.include_router(router, prefix="/api")
    with TestClient(app) as client:
        response = client.post(f"/api/documents/{old['document']['id']}/reimport")
        assert response.status_code == 200, response.text
        assert response.json()["version"]["version"] == 2
        assert response.json()["document"]["id"] == old["document"]["id"]
        assert client.post("/api/documents/not-present/reimport").status_code == 404
        Path(response.json()["version"]["path"]).write_bytes(b"tampered fixture")
        invalid = client.post(f"/api/documents/{old['document']['id']}/reimport")
        assert invalid.status_code == 422
        assert "immutable" in invalid.text
        assert db.one("SELECT COUNT(*) AS n FROM document_versions")["n"] == 2
